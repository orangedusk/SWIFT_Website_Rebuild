"""Put the client's edits from the copy deck back into the website.

    python3 tools/apply_copy.py SWIFT_copy_deck.xlsx            # preview: lists what would change
    python3 tools/apply_copy.py SWIFT_copy_deck.xlsx --apply    # make the changes and rebuild the pages

Website copy sheet: every row with something in 'New copy' replaces its 'Current copy' wherever it
appears in the page sources (index.html, request-appointment.html, tools/build_*.py). Shared text
(menus, footer) is in both index.html and request-appointment.html, and both are updated.
Symptom search sheet: every row with a 'New outcome' changes that term's outcome in symptoms.js.

Rows it can't place (for example text split around a link) are listed as 'Needs a manual edit'.
Review the preview, then commit the result on the demo branch as usual.
"""
import os, re, sys, html, subprocess
from openpyxl import load_workbook

ROOT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), '')
SOURCES = ['index.html', 'request-appointment.html', 'tools/build_fees.py', 'tools/build_services.py',
           'tools/build_team.py', 'tools/build_gallery.py', 'tools/build_faq.py', 'tools/build_page.py']
OUTCOME_CODES = {'Call 000': '000', 'Call 000 or Poisons Information (13 11 26)': 'poisons',
                 'Not a SWIFT walk-in (call SWIFT; 000 in an emergency)': 'out-of-scope',
                 'Call SWIFT first': 'call-first', 'Come to SWIFT': 'swift', 'On-site dentist': 'dental',
                 'See your GP': 'gp'}


def pattern_for(text):
    """Regex that finds the copy however the source writes it: entities, curly or straight
    apostrophes, escaped quotes, and line breaks between words."""
    parts = []
    for ch in text:
        if ch.isspace(): parts.append(r'(?:\s|&nbsp;)+'); continue
        if ch in "'’": parts.append(r"(?:'|’|&rsquo;|&#39;|\\')"); continue
        if ch == '"': parts.append(r'(?:"|&quot;|\\")'); continue
        if ch == '&': parts.append(r'(?:&amp;|&)'); continue
        if ch == '“': parts.append(r'(?:“|&ldquo;)'); continue
        if ch == '”': parts.append(r'(?:”|&rdquo;)'); continue
        if ch == '—': parts.append(r'(?:—|&mdash;)'); continue
        if ch == '–': parts.append(r'(?:–|&ndash;)'); continue
        parts.append(re.escape(ch))
    # Whole piece of text only: it must fill a whole element (between > and <) or a whole string
    # (between quotes), so 'Where to go' never matches inside 'Where to go, from most to least urgent'.
    body = re.sub(r'(\(\?:\\s\|&nbsp;\)\+)+', r'(?:\\s|&nbsp;)+', ''.join(parts))
    return re.compile(r'(>|(?<!\\)[\'"])(\s*)' + body + r'(\s*)(?=<|[\'"])')


def quote_context(src, pos):
    """If pos sits inside a JS/Python string on its line, return its quote character."""
    line_start = src.rfind('\n', 0, pos) + 1
    seg, q, i = src[line_start:pos], None, 0
    while i < len(seg):
        c = seg[i]
        if c == '\\': i += 2; continue
        if q is None and c in '\'"': q = c
        elif q == c: q = None
        i += 1
    return q


def encode(new, src, pos):
    out = html.escape(new, quote=False).replace('’', '&rsquo;')
    q = quote_context(src, pos)
    if q: out = out.replace('\\', '\\\\').replace(q, '\\' + q)
    return out


def main(path, apply):
    wb = load_workbook(path, data_only=True)
    files = {f: open(ROOT + f).read() for f in SOURCES}
    changed, manual = [], []

    ws = wb['Website copy']
    edits = {}   # current copy -> (ids, new copy); the same text in several rows is changed once
    for row in ws.iter_rows(min_row=2, values_only=True):
        rid, page, section, element, current, new, notes = (list(row) + [None] * 7)[:7]
        if not rid or str(rid).startswith('EXAMPLE') or not new or not str(new).strip(): continue
        current, new = str(current).strip(), str(new).strip()
        if new == current: continue
        if current in edits and edits[current][1] != new:
            manual.append((rid, current, new + '   [conflicts with ' + edits[current][0][0] + ']', 0)); continue
        edits.setdefault(current, ([], new))[0].append(rid)
    for current, (ids, new) in edits.items():
        pat, hits = pattern_for(current), 0
        for f in SOURCES:
            src, out, last = files[f], [], 0
            for m in pat.finditer(src):
                out.append(src[last:m.start()]); out.append(m.group(1) + m.group(2) + encode(new, src, m.end(2)) + m.group(3))
                last = m.end(); hits += 1
            if out: files[f] = ''.join(out) + src[last:]
        (changed if hits else manual).append((', '.join(ids), current, new, hits))

    sym_changes = []
    if 'Symptom search' in wb.sheetnames:
        js = open(ROOT + 'symptoms.js').read()
        for row in wb['Symptom search'].iter_rows(min_row=2, values_only=True):
            term, cat, cur, item, new, notes = (list(row) + [None] * 6)[:6]
            if not term or not new or new == cur: continue
            code = OUTCOME_CODES.get(str(new).strip())
            if not code: manual.append(('search: ' + term, cur, new, 0)); continue
            pat = re.compile(r"(\[\s*(['\"])%s\2,\s*)'[^']+'" % re.escape(term))
            js, n = pat.subn(lambda m: m.group(1) + "'%s'" % code, js)
            (sym_changes if n else manual).append(('search: ' + term, cur, new, n))
        if apply and sym_changes: open(ROOT + 'symptoms.js', 'w').write(js)

    print('Will change:' if not apply else 'Changed:')
    for rid, cur, new, n in changed + sym_changes:
        print('  %-40s %d place(s)\n      was: %s\n      now: %s' % (rid, n, cur, new))
    if manual:
        print('\nNeeds a manual edit (text not found as one phrase, often because it wraps a link):')
        for rid, cur, new, n in manual: print('  %-40s %s  ->  %s' % (rid, cur, new))
    if not (changed or sym_changes or manual): print('  nothing: no New copy or New outcome filled in')

    if apply and changed:
        for f in SOURCES: open(ROOT + f, 'w').write(files[f])
        for s in ('fees', 'team', 'services', 'gallery', 'faq'):
            subprocess.run([sys.executable, ROOT + 'tools/build_%s.py' % s], check=True, stdout=subprocess.DEVNULL)
        print('\nPages rebuilt. Check the site, then commit on the demo branch.')
    elif not apply and (changed or sym_changes):
        print('\nPreview only. Run again with --apply to make these changes.')


if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if not args: sys.exit(__doc__)
    main(args[0], '--apply' in sys.argv)
