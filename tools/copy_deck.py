"""Export every piece of website copy into an Excel copy deck for client review.

    python3 tools/copy_deck.py            # writes SWIFT_copy_deck.xlsx in the project root

Sheets:
  Read me         how to fill it in
  Website copy    one row per piece of text: ID, Page, Section, Element, Current copy, New copy, Notes
  Symptom search  every term in symptoms.js and where it sends the patient

Shared parts (emergency banner, header, menus, footer) are listed once, under 'All pages'.
Demo-only tools (tour, campaign examples, password screen, sample profiles) are not included.
Apply the client's edits with tools/apply_copy.py.
"""
import os, re, html
from html.parser import HTMLParser

ROOT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), '')
PAGES = [('index.html', 'Home'), ('services.html', 'Services'), ('fees.html', 'Fees'),
         ('team.html', 'Team'), ('gallery.html', 'Gallery'), ('request-appointment.html', 'Request appointment')]

# Elements that hold one piece of copy each
LEAF = {'h1', 'h2', 'h3', 'h4', 'p', 'li', 'a', 'button', 'summary', 'label', 'dt', 'dd',
        'figcaption', 'th', 'td', 'option', 'span', 'legend', 'caption'}
INLINE = {'a', 'strong', 'em', 'b', 'i', 'br', 'abbr', 'small', 'code'}
SKIP = {'script', 'style', 'svg', 'noscript', 'template', 'dialog'}
VOID = {'br', 'img', 'input', 'meta', 'link', 'hr', 'source', 'area', 'col', 'embed', 'wbr'}
NAMES = {'h1': 'Page heading', 'h2': 'Heading', 'h3': 'Sub-heading', 'h4': 'Small heading',
         'p': 'Text', 'li': 'List item', 'a': 'Link / button', 'button': 'Button', 'summary': 'Question',
         'label': 'Form label', 'dt': 'Label', 'dd': 'Value', 'figcaption': 'Caption', 'th': 'Table label',
         'td': 'Table value', 'option': 'Option', 'span': 'Text', 'legend': 'Form heading', 'caption': 'Table caption'}


class Node:
    def __init__(self, tag, attrs, parent):
        self.tag, self.attrs, self.parent, self.children = tag, dict(attrs), parent, []


class Tree(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node('root', [], None); self.cur = self.root
    def handle_starttag(self, tag, attrs):
        n = Node(tag, attrs, self.cur); self.cur.children.append(n)
        if tag not in VOID: self.cur = n
    def handle_startendtag(self, tag, attrs):
        self.cur.children.append(Node(tag, attrs, self.cur))
    def handle_endtag(self, tag):
        n = self.cur
        while n is not self.root and n.tag != tag: n = n.parent
        if n is not self.root: self.cur = n.parent
    def handle_data(self, data):
        self.cur.children.append(data)


def text_of(n):
    if isinstance(n, str): return n
    if n.tag in SKIP: return ''
    if n.tag == 'br': return ' '
    if 'sr-only' in n.attrs.get('class', '') and n.tag != 'h2': return ''
    return ''.join(text_of(c) for c in n.children)

def clean(t): return re.sub(r'\s+', ' ', t).strip()

def has_leaf_child(n):
    for c in n.children:
        if isinstance(c, Node) and c.tag not in SKIP:
            if c.tag in LEAF and c.tag not in INLINE and clean(text_of(c)): return True
            if c.tag == 'span' and 'block' in c.attrs.get('class', '').split() and clean(text_of(c)): return True
            if c.tag not in INLINE and has_leaf_child(c): return True
    return False

def own_inline_links(n):
    return any(isinstance(c, Node) and c.tag == 'a' for c in n.children)


def walk(n, section, out, state):
    """state['h'] holds the most recent section heading, used as the Section for each row."""
    if isinstance(n, str) or n.tag in SKIP: return
    if 'hidden' in n.attrs and n.tag in ('section', 'div'): return          # switched-off content
    if n.attrs.get('id') == 'bookingEmbed': return                           # not live until online booking is connected
    if n.attrs.get('id') == 'visitTileContent': return                       # filled from the steps list (see visit_steps)
    if n.attrs.get('aria-hidden') == 'true': return
    if n.tag in ('h1', 'h2', 'legend') and section is None:
        t = clean(text_of(n))
        if t: state['h'] = t
    for attr, label in (('placeholder', 'Placeholder'), ('alt', 'Image description')):
        v = n.attrs.get(attr)
        if v and clean(v): out.append((section or state['h'], label, clean(v)))
    t = clean(text_of(n))
    is_leaf_tag = n.tag in LEAF and not (n.tag == 'span' and not t)
    if is_leaf_tag and t and not has_leaf_child(n):
        if n.tag in INLINE and n.parent is not None and n.parent.tag in LEAF and clean(text_of(n.parent)) != t:
            return
        el = NAMES.get(n.tag, 'Text')
        if n.tag in ('p', 'li', 'span', 'dd') and own_inline_links(n): el += ' (contains a link)'
        if n.tag in ('h1', 'h2', 'legend') and section is None: el = 'Page heading' if n.tag == 'h1' else 'Section heading'
        out.append((section or state['h'], el, t)); return
    for c in n.children: walk(c, section, out, state)


def find(n, pred):
    if isinstance(n, str): return None
    if pred(n): return n
    for c in n.children:
        r = find(c, pred)
        if r: return r
    return None


def head_rows(tree):
    rows = []
    t = find(tree.root, lambda n: n.tag == 'title')
    if t: rows.append(('Page title (browser tab)', 'Title', clean(text_of(t))))
    for name, label in (('description', 'Search engine description'),):
        m = find(tree.root, lambda n: n.tag == 'meta' and n.attrs.get('name') == name)
        if m: rows.append(('Search results', label, clean(m.attrs.get('content', ''))))
    m = find(tree.root, lambda n: n.tag == 'meta' and n.attrs.get('property') == 'og:description')
    if m: rows.append(('Link previews', 'Shared-link description', clean(m.attrs.get('content', ''))))
    return rows


# JavaScript copy: strings in the page scripts that people read (steps, search results, form messages)
def js_strings(code):
    """Yield string literals from JS source, skipping comments and regex literals."""
    i, n, prev = 0, len(code), ''
    while i < n:
        c = code[i]
        if code.startswith('//', i):
            i = code.find('\n', i); i = n if i < 0 else i; continue
        if code.startswith('/*', i):
            i = code.find('*/', i + 2); i = n if i < 0 else i + 2; continue
        if c in '\'"`':
            j, buf = i + 1, []
            while j < n and code[j] != c:
                if code[j] == '\\' and j + 1 < n: buf.append(code[j + 1]); j += 2; continue
                buf.append(code[j]); j += 1
            yield ''.join(buf); i = j + 1; prev = c; continue
        if c == '/' and prev in '(,=:[!&|?{};+' :
            j = i + 1
            while j < n and code[j] not in '/\n':
                j += 2 if code[j] == '\\' else 1
            i = j + 1; prev = '/'; continue
        if not c.isspace(): prev = c
        i += 1

def looks_like_copy(raw):
    s = re.sub(r'^[^<>]*">', '', raw)                     # drop a tail of an HTML attribute
    plain = clean(html.unescape(re.sub(r'<[^>]+>', ' ', s)))
    words = plain.split()
    if sum(1 for w in words if re.search(r'[A-Za-z]', w)) < 2 or not re.search(r'[A-Za-z]{3}', plain): return None
    if re.search(r'(https?://|\.html|\.jpe?g|\.png|brand_assets|tel:|=>|\bfunction\b|\bvar\b)', raw): return None
    techy = sum(1 for w in words if re.search(r'[\[\]=:/{}()#]|^-|-$|^[a-z]+-[a-z0-9-]+$', w))
    if techy / len(words) > 0.25: return None
    if re.fullmatch(r'[a-z0-9 _.,-]+', plain) and len(words) < 4: return None
    if plain.startswith('"') or plain in ('IBM Plex Sans',): return None
    return plain

STEP_FIELDS = (('title', 'Step title'), ('desc', 'Step text'), ('tip', 'Good to know'))
FIELD = r"""%s:\s*(?:'((?:[^'\\]|\\.)*)'|"((?:[^"\\]|\\.)*)")"""

def visit_steps(src):
    """The seven visit steps live in a JS list on the homepage; give each field its own row."""
    m = re.search(r'var steps = \[(.*?)\n    \];', src, re.S)
    if not m: return [], ''
    rows = []
    for i, step in enumerate(re.findall(r'\{(.*?)\}', m.group(1), re.S), 1):
        for key, label in STEP_FIELDS:
            f = re.search(FIELD % key, step)
            if f: rows.append(('What happens when you visit', 'Step %d: %s' % (i, label), (f.group(1) or f.group(2)).replace("\\'", "'")))
    return rows, m.group(0)

def js_rows(src):
    steps, steps_src = visit_steps(src)
    rows = list(steps)
    for block in re.findall(r'<script>(.*?)</script>', src, re.S):
        if steps_src: block = block.replace(steps_src, '')
        for raw in js_strings(block):
            for piece in re.split(r'<[^>]*>|<[a-z][^>]*$', raw):      # one row per piece of text between tags
                if '="' in piece or piece.lstrip().startswith('"'): continue
                plain = looks_like_copy(piece)
                if plain: rows.append(('Messages shown by the page', 'Message', plain))
    return rows


def collect():
    rows = []
    for fname, page in PAGES:
        src = open(ROOT + fname).read()
        tree = Tree(); tree.feed(src)
        body = find(tree.root, lambda n: n.tag == 'body')
        for sec, el, t in head_rows(tree): rows.append((page, sec, el, t))
        if fname == 'index.html':
            # site chrome, once
            for n in body.children:
                if isinstance(n, Node) and n.tag in ('div', 'header', 'nav', 'footer') and n.tag != 'main':
                    label = {'header': 'Header and menu', 'nav': 'Mobile bottom bar', 'footer': 'Footer'}.get(n.tag, 'Emergency banner')
                    if n.attrs.get('id') == 'top': continue
                    out = []; walk(n, label, out, {'h': label})
                    for sec, el, t in out: rows.append(('All pages', sec if sec in ('Header and menu', 'Mobile bottom bar', 'Footer', 'Emergency banner') else label, el, t))
        main = find(body, lambda n: n.tag == 'main')
        out = []; walk(main, None, out, {'h': 'Top of page'})
        for sec, el, t in out: rows.append((page, sec, el, t))
        dialog = find(body, lambda n: n.tag == 'dialog')
        if fname == 'request-appointment.html':
            pass
        for sec, el, t in js_rows(src):
            rows.append((page, sec, el, t))
    return rows


SECTION_NAMES = {
    'what-we-do': 'What we do', 'where-to-go': 'Where to go (three columns)',
    'quickRoutesHeading': 'Quick-link tiles', 'visit': 'What happens when you visit', 'services': 'Our clinical team treats',
    'other-care': 'Here for something else?', 'fees': 'Fees, up front', 'doctors': 'Our doctors',
    'facilities': 'Our clinic facilities', 'location': 'Find us', 'faq': 'Common questions',
    'whatWeDoHeading': 'What we do', 'doctorsHeading': 'Our doctors',
    'facilitiesHeading': 'Our clinic facilities', 'otherCareHeading': 'Here for something else?',
    'medicare': 'How you pay', 'pathology': 'Pathology', 'emergency': 'Emergency & urgent care',
    'infusion': 'Infusions', 'wound': 'Wound care', 'radiology': 'Scans and imaging', 'extras': 'Other costs',
    'grp0': 'Our Doctors', 'grp1': 'Our Nurses', 'grp2': 'Our Advisors',
}


def build_rows():
    raw = collect()
    rows, seen, counters = [], set(), {}
    last_section = {}
    for page, sec, el, t in raw:
        if not t or len(t) < 2: continue
        key = (page, t, el)
        if key in seen: continue
        seen.add(key)
        section = SECTION_NAMES.get(sec, sec)
        if section and section[0].islower() and '-' in section: section = section.replace('-', ' ').capitalize()
        pslug = re.sub(r'[^a-z]+', '-', page.lower()).strip('-')
        sslug = '-'.join(re.sub(r'[^a-z0-9 ]+', '', section.lower()).split()[:3])
        counters[(pslug, sslug)] = counters.get((pslug, sslug), 0) + 1
        rows.append(('%s.%s.%02d' % (pslug, sslug, counters[(pslug, sslug)]), page, section, el, t))
    return rows


OUTCOMES = {'000': 'Call 000', 'poisons': 'Call 000 or Poisons Information (13 11 26)',
            'out-of-scope': 'Not a SWIFT walk-in (call SWIFT; 000 in an emergency)',
            'call-first': 'Call SWIFT first', 'swift': 'Come to SWIFT', 'dental': 'On-site dentist',
            'gp': 'See your GP'}


def symptom_rows():
    src = open(ROOT + 'symptoms.js').read()
    rows, cat = [], ''
    for line in src.split('\n'):
        c = re.match(r'\s*//\s*(.+)', line)
        if c and 'INTERIM' not in line and not line.strip().startswith('// -'): cat = c.group(1).strip()
        m = re.match(r"\s*\[(\"[^\"]*\"|'[^']*'),\s*'([^']+)',\s*(null|'[^']*')(?:,\s*'([^']*)')?\]", line)
        if m:
            term = m.group(1)[1:-1]
            rows.append((term, cat, OUTCOMES.get(m.group(2), m.group(2)), '' if m.group(3) == 'null' else m.group(3).strip("'")))
    return rows


def write(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    F = 'Arial'
    hdr_font = Font(name=F, bold=True, color='FFFFFF'); hdr_fill = PatternFill('solid', fgColor='1E453F')
    body = Font(name=F, size=10); grey = Font(name=F, size=10, color='555555')
    input_fill = PatternFill('solid', fgColor='FFF7CC'); example_font = Font(name=F, size=10, italic=True, color='7A6A00')
    thin = Side(style='thin', color='D9E3E0'); border = Border(bottom=thin)
    wrap = Alignment(wrap_text=True, vertical='top')

    # Read me
    ws = wb.active; ws.title = 'Read me'
    lines = [
        ('SWIFT website copy deck', Font(name=F, bold=True, size=14)),
        ('Every piece of text on the new SWIFT website, for review and sign-off.', body),
        ('', body),
        ('How to review', Font(name=F, bold=True, size=11)),
        ('1. Go to the "Website copy" sheet. Each row is one piece of text, grouped by page and section.', body),
        ('2. To change text, type the new wording in the yellow "New copy" column. Leave it blank if the current copy is fine.', body),
        ('3. Use "Notes" for questions or comments. Rows with a note but no new copy are treated as questions, not changes.', body),
        ('4. In the "Symptom search" sheet, change where a search sends the patient by picking from the yellow "New outcome" list. Clinical routing is a draft until Dr Manivel signs it off.', body),
        ('5. Don\'t change the ID, Page, Section or Current copy columns. They are how the new wording is put back into the website.', body),
        ('', body),
        ('Good to know', Font(name=F, bold=True, size=11)),
        ('"All pages" rows (banner, header, menus, footer) appear on every page, so one change updates them everywhere.', body),
        ('Rows marked "(contains a link)" include linked words. Keep the linked words in your new copy, or note what the link should say.', body),
        ('"Message (shown by the page)" rows appear only after an action, for example a search result or a form error.', body),
        ('The business name "SWIFT Emergency & Urgent Care" is kept as is (client instruction).', body),
        ('Demo-only content (the guided tour, campaign examples, sample doctor profiles) is not included.', body),
    ]
    for i, (t, f) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 120

    # Website copy
    ws = wb.create_sheet('Website copy')
    heads = ['ID', 'Page', 'Section', 'Element', 'Current copy', 'New copy', 'Notes']
    widths = [30, 14, 30, 24, 70, 70, 40]
    for i, (h, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(row=1, column=i, value=h); c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(vertical='center')
        ws.column_dimensions[c.column_letter].width = w
    rows = build_rows()
    # Example row: shows the format; starts with EXAMPLE so tools/apply_copy.py skips it
    ex = ['EXAMPLE (ignored)', 'Home', 'Common questions', 'Question', 'What ages does SWIFT treat?',
          'Which ages does SWIFT see?', 'Example only: shows how to suggest new wording']
    for j, v in enumerate(ex, 1):
        c = ws.cell(row=2, column=j, value=v); c.font = example_font; c.alignment = wrap
        if j == 6: c.fill = input_fill
    for r, row in enumerate(rows, 3):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r, column=j, value=v); c.font = grey if j in (1, 4) else body; c.alignment = wrap; c.border = border
        for j in (6, 7):
            c = ws.cell(row=r, column=j); c.font = body; c.alignment = wrap; c.border = border
            if j == 6: c.fill = input_fill
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = 'A1:G%d' % (len(rows) + 2)

    # Symptom search
    ws = wb.create_sheet('Symptom search')
    heads = ['Search term', 'Category', 'Current outcome', 'Highlights in column', 'New outcome', 'Notes']
    widths = [30, 40, 44, 20, 44, 40]
    for i, (h, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(row=1, column=i, value=h); c.font = hdr_font; c.fill = hdr_fill
        ws.column_dimensions[c.column_letter].width = w
    srows = symptom_rows()
    for r, row in enumerate(srows, 2):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r, column=j, value=v); c.font = body; c.alignment = wrap; c.border = border
        for j in (5, 6):
            c = ws.cell(row=r, column=j); c.font = body; c.border = border
            if j == 5: c.fill = input_fill
    dv = DataValidation(type='list', formula1='"%s"' % ','.join(v.replace(',', ';') for v in OUTCOMES.values()), allow_blank=True)
    dv.error = 'Pick an outcome from the list'; ws.add_data_validation(dv); dv.add('E2:E%d' % (len(srows) + 1))
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = 'A1:F%d' % (len(srows) + 1)

    wb.save(path)
    return len(rows), len(srows)


if __name__ == '__main__':
    out = ROOT + 'SWIFT_copy_deck.xlsx'
    n, m = write(out)
    print('wrote', out, '|', n, 'copy rows,', m, 'search terms')
